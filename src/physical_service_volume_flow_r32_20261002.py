"""DEA-anchored finite-volume/hydraulic NECESSARY service feasibility screens.

Not a sufficient hardware design: full hot composite curves and turbine states
are not published for all configurations. Failure in the optimistic thermal
upper box rejects service delivery. Passing it is not certified attainability.
DEA public July2026 workbook supplies finite volume, rate, temperature and pumps.
Daily losses are % of rated capacity in source: a SOC-proportional decay is an
optimistic structural approximation, not exact tank cooling replication.
"""
from pathlib import Path
import csv,json,importlib.util,os
from datetime import datetime,timezone
import numpy as np
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('service_lp',ROOT/'src/physical_service_feasibility_r32_20261002.py')
m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
RAW=os.environ.get('PROJECT37_STORAGE_WORKBOOK')

def stores():
    # The attributed numerical subset avoids requiring redistribution of the
    # original official workbook. Extraction/revalidation is a separate step.
    subset=ROOT/'datasets/dea_storage_141_transcription_r32_20261002.json'
    if subset.exists():
        return json.loads(subset.read_text(encoding='utf-8'))['records']
    if not RAW:
        raise FileNotFoundError('Missing released numerical subset. For explicit original-source revalidation, set PROJECT37_STORAGE_WORKBOOK to an accessible workbook; no local raw path is assumed.')
    book=load_workbook(Path(RAW),data_only=True,read_only=True)
    records=[dict(name='no_store',volume_m3=0,hot_C=99,service_C=97,rate_MW=0,loss_daily=0,aux=0,nominal_capacity_MWh=0,nominal_capex_EUR2025=0,fixed_OM_EUR_y=0)]
    for name,volrow,temprow,disrow,fixedrow in [('141a TTES',28,29,30,25),('141b Large TTES',29,30,31,26)]:
        s=book[name]
        def val(row):return float(s.cell(row,3).value)
        cap=val(8)
        records.append(dict(name=name,volume_m3=val(volrow),hot_C=val(temprow),service_C=val(disrow),rate_MW=val(9),
                            loss_daily=val(14)/100,aux=val(15)/100,nominal_capacity_MWh=cap,
                            nominal_capex_EUR2025=val(21)*(cap/1000)*1e6,fixed_OM_EUR_y=val(fixedrow)*cap,
                            capex_scope='Catalogue Danish reference, not project quote;5000m3 PDF excludes external connection while workbook note includes some connection: preserve scope conflict.50000m3 includes connection.'))
    return records

def main():
    rows=[];templates=stores()
    # This is a commissioning-gate pilot, not an annual case/demand-year matrix.
    for profile,demand in [('constant',np.full(168,36.5)),('declared_synthetic_week',36.5*(1+.22*np.cos(np.arange(168)/168*2*np.pi)))]:
        for ts in [80,100,120]:
            #150/450kg/s are designed service contracts, not observed networks.
            for flow in [150,450]:
                for s in templates:
                    # Fixed volume, temperature-difference dependent useful energy.
                    cap=.9*1000*s['volume_m3']*.00418*(s['hot_C']-60)/3600
                    for cascade in [False,True]:
                        for capture in [False,True]:
                            r=m.solve(demand,ts,60,capture,cascade,0,
                                      tank_hot_C=s['hot_C'],tank_approach_K=s['hot_C']-s['service_C'],
                                      usable_capacity_MWh=cap,store_power_MW=s['rate_MW'],flow_limit_kg_s=flow,
                                      weekly_loss=1-(1-s['loss_daily'])**7,charge_efficiency=1,discharge_efficiency=1,
                                      steam_grade_override_C=390,compressor_grade_override_C=158,
                                      auxiliary_fraction=s['aux'])
                            r.update(profile=profile,tank_type=s['name'],volume_m3=s['volume_m3'],
                                     required_nominal_flow_kg_s=36.5/(.00418*(ts-60)),
                                     nominal_tank_capex_catalogue_EUR2025=s['nominal_capex_EUR2025'],
                                     nominal_fixed_OM_catalogue_EUR2025_y=s['fixed_OM_EUR_y'],
                                     feasibility_interpretation='NECESSARY optimistic screen: steam treated as freely high-grade up to main400C-pinch; compressor total treated at maximum hot173C-pinch15K. Passing does not prove physical attainability.',
                                     costs_interpretation='Reference tank cost only. No total CCS/network investment rank; no foreign-market quotation; no added variable O&M charge because it already contains pump electricity.')
                            rows.append(r)
    with (ROOT/'results/physical_service_volume_flow_r32_20261002.csv').open('w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
    def get(profile,ts,flow,store,capture,model):
        return next(r for r in rows if r['profile']==profile and r['supply_temperature_C']==ts and r['hydraulic_flow_limit_kg_s']==flow and r['tank_type']==store and r['capture_on']==capture and r['model']==model)
    comparisons=[]
    for profile in ['constant','declared_synthetic_week']:
        for ts in [80,100,120]:
            for flow in [150,450]:
                for s in templates:
                    base=get(profile,ts,flow,s['name'],False,'temperature_cascade')
                    c=get(profile,ts,flow,s['name'],True,'temperature_cascade')
                    approx=get(profile,ts,flow,s['name'],True,'aggregate_heat')
                    comparisons.append(dict(profile=profile,ts=ts,flow=flow,store=s['name'],
                                            paired_necessary_gate_pass=bool(base['service_feasible'] and c['service_feasible']),
                                            noCCS_gate_pass=base['service_feasible'],CCS_gate_pass=c['service_feasible'],
                                            aggregate_CCS_gate_pass=approx['service_feasible'],unserved_CCS_MWh=c['unserved_MWh'],
                                            aggregate_false_pass=bool(base['service_feasible'] and approx['service_feasible'] and not c['service_feasible'])))
    flags=[r for r in comparisons if r['aggregate_false_pass']]
    report=dict(created_utc=datetime.now(timezone.utc).isoformat(),LPs=len(rows),pairs=len(comparisons),
                aggregate_false_pass_count=len(flags),aggregate_false_pass_cases=flags,
                tank_templates=templates,
                nominal_flow_requirements_kg_s={str(ts):36.5/(.00418*(ts-60)) for ts in [80,100,120]},
                max_eq_error=max(r['max_eq_error'] for r in rows),max_ub_violation=max(r['max_ub_violation'] for r in rows),
                total_solver_seconds=sum(r['solver_seconds'] for r in rows),
                gate='Actual finite-volume/rate/flow screen computed. NOT full two-case physical validation or cost-optimal investment rank.',
                source_version='DEA storage catalogue Chapter141 revised July2026; new official XLSX downloaded2026-10-02;2025 reference conditions. CCS catalogue remains April2024/February2024.',
                source_conflicts=['PDF max water about97C versus workbook charging99C/discharged97C: retain both; use99/97 as optimistic charging/discharge bound.',
                                 'Workbook daily heat loss is % rated capacity, not % state. SOC-proportional version is an optimistic necessary screen.',
                                 '5000m3 PDF excludes external network pumps/HX/connection, workbook noteJ partly includes connection. Cannot use it as complete installed quote.',
                                 'Variable O&M includes pumping electricity: do not add it to separately calculated electricity cost.',
                                 'Round-trip98% is annual including standing losses: do not multiply another0.98 on top of charge/discharge100% and standing loss.'])
    (ROOT/'results/physical_service_volume_flow_qc_r32_20261002.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report,indent=2))

if __name__=='__main__':main()
